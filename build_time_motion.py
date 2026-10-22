#!/usr/bin/env python3
"""Render data/qaqc-time-motion.xlsx into the pilot's time and motion pages:
one page per project (Lagoons - Ibiza at index.html, Hills - Gems beside it).

The workbook has one tab per unit. Down each tab, one row per activity with its
minutes and its clock times (HHMM, so 957 is 09:57). Each fill colour is a
different person: the "(total)" row of a colour is that person's activity, and
the rows below it in the same colour are its sub-tasks. This page lays each
unit out by person and activity, sub-tasks in a dropdown, with whatever the
sub-tasks do not account for shown as "Other functional". A tab lands on the
page whose project its name starts with (see SITES).

As elsewhere on the site, nothing is corrected. A blank row is "Activity
pending"; 0 minutes or "not recorded seperately" is "Not recorded separately";
and minutes that disagree with the row's own clock times by more than rounding
are flagged on the page, not fixed.

Needs openpyxl (pip install openpyxl).
"""

import csv
import math
import re
from pathlib import Path

import theme
from theme import GENERATED, esc

ROOT = Path(__file__).parent
XLSX_PATH = ROOT / "data" / "qaqc-time-motion.xlsx"
# Floor area and bedrooms per unit; the workbook does not carry them.
UNITS_PATH = ROOT / "data" / "time-motion-units.csv"

# Floor-area normalisation: recorded time per this many square metres.
PER_AREA = 100

# One page per project: (theme.PAGES key, workbook tab prefix, page heading).
# The tab label and file come from theme.PAGES.
SITES = [
    ("tm-ibiza", "Ibiza", "Lagoons \u2013 Ibiza"),
    ("tm-gems", "Gems", "Hills \u2013 Gems"),
]

# Tabs that are not units on any page.
EXCLUDE = {"Gems-E511B_old"}  # superseded by Gems-E511B

# Units read from the workbook but not shown (another unit may borrow from
# them), and not mentioned on the page.
HIDDEN = {"Ibiza T244"}

# A unit whose capture survey is the one logged on another tab:
# {unit: tab it is logged on}. Per the pilot team, T242 was surveyed with
# Matterport at the time logged on T244's tab, not with iGUIDE.
CAPTURE_FROM = {"Ibiza T242": "Ibiza T244"}

# Minutes within this of the row's own clock times are rounding, not flagged.
ROUNDING = 2

# Each fill colour in the sheet is one person, keyed by the Office theme slot
# the sheet fills with. The page colours keep the sheet's hues (orange, blue,
# purple, green) as a categorical set validated for colour-blind separation.
PEOPLE = {
    5: ("mepg", "MEP inspector"),
    4: ("meps", "MEP specialist"),
    8: ("qa", "QA inspector"),
    9: ("cap", "Capture operator"),
}
SLOTS = [slot for slot, _ in PEOPLE.values()]

UNLOGGED = "Start and end time not logged"

# Minutes for an activity the sheet leaves blank, by person. They count towards
# the unit's totals, and the row says its times were not logged.
UNTIMED = {"meps": 30}

# Tests the sheet does not carry. part_of is the person whose activity a test
# belongs to: it is added there as a sub-task and counts towards their time.
# Without one (the flood test is the contractor's) it is shown but kept out of
# the totals. A test has either flat minutes, for every unit or those in units,
# or a logged start and end per unit, which give its minutes (man-hours) and
# draw it on the unit's timeline with the activity it is part of.
TESTS = [
    dict(name="Flood test", who="Contractor", minutes=30, status="Completed",
         note="Conducted by contractor, Verified by functional QA.", part_of=None),
    dict(name="Static test", who="Functional QA", minutes=20, status="Completed",
         note="Conducted by functional QA.", part_of="qa"),
    dict(name="Drain check", who="MEP team", status="Completed",
         note="Conducted by MEP team.", part_of="mepg",
         logged={"Ibiza T240": ("11:21", "12:50"),     # logged on site
                 "Ibiza T241": ("12:52", "13:52"),
                 "Ibiza T242": ("13:57", "15:08")}),
]


def applies(test, u):
    if "logged" in test:
        return u["name"] in test["logged"]
    return not test.get("units") or u["name"] in test["units"]


def test_time(test, u):
    """(minutes, (start, end) in minutes past midnight, or None) for a test on this unit."""
    if "logged" in test:
        start, end = (int(h) * 60 + int(m) for h, m in
                      (t.split(":") for t in test["logged"][u["name"]]))
        return end - start, (start, end)
    return test["minutes"], None


def and_join(xs):
    xs = list(xs)
    return " and ".join(xs) if len(xs) < 3 else ", ".join(xs[:-1]) + " and " + xs[-1]


def short_unit(name):
    """'Ibiza T241' as 'T241', for lists of units on a page that is all one project."""
    return re.sub(rf"^(?:{'|'.join(re.escape(p) for _, p, _ in SITES)})[\s-]*", "", name)

STATE_TEXT = {"pending": "Activity pending", "separate": "Not recorded separately"}


# ---------------------------------------------------------------- parsing ---

def fmt_dur(mins):
    """Hours and minutes, as on the activity log."""
    h, m = divmod(mins, 60)
    return f"{h}h {m:02d}m" if h else f"{m}m"


def fmt_clock(t):
    return f"{t // 60:02d}:{t % 60:02d}"


def fmt_windows(sessions, sep=", "):
    return sep.join(f"{fmt_clock(s)}&ndash;{fmt_clock(e)}" for s, e in sessions)


def plain(s):
    """Tooltip text: the entities fmt_windows emits, as characters."""
    return s.replace("&ndash;", "–")


def clock(v):
    """Minutes past midnight from an HHMM integer. Blank and 0 are no time."""
    if not isinstance(v, (int, float)) or not v:
        return None
    h, m = divmod(int(v), 100)
    return h * 60 + m if h < 24 and m < 60 else None


def reading(v):
    """What the minutes cell says: (state, minutes)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return "pending", None
    if isinstance(v, (int, float)):
        return ("separate", None) if v == 0 else ("recorded", round(v))
    if "not recorded" in v.lower():
        return "separate", None
    return "text", v.strip()


def fill_key(cell):
    f = cell.fill
    if not f or not f.fill_type:
        return None
    c = f.fgColor
    return c.theme if c.type == "theme" else c.rgb


def session_pairs(ws):
    """Columns of start/end pairs, from the header row. The sheet runs a
    second pair along where an activity was split over two sessions."""
    def head(col):
        return re.sub(r"\s*\d+$", "", str(ws.cell(1, col).value or "").strip().lower())
    pairs, c = [], 3
    while (head(c), head(c + 1)) == ("start", "end"):
        pairs.append((c, c + 1))
        c += 2
    return pairs


def read_row(ws, r, pairs):
    state, minutes = reading(ws.cell(r, 2).value)
    sessions, flags = [], []
    for cs, ce in pairs:
        s, e = clock(ws.cell(r, cs).value), clock(ws.cell(r, ce).value)
        if s is None or e is None:
            continue
        if e <= s:
            flags.append(f"Ends {fmt_clock(e)}, before its start at {fmt_clock(s)}.")
            continue
        sessions.append((s, e))
    sessions.sort()
    if state == "recorded" and sessions:
        span = sum(e - s for s, e in sessions)
        if abs(span - minutes) > ROUNDING:
            flags.append(f"{plain(fmt_windows(sessions))} is {fmt_dur(span)}; "
                         f"the sheet records {fmt_dur(minutes)}.")
    return {"state": state, "minutes": minutes if state == "recorded" else None,
            "text": minutes if state == "text" else None,
            "sessions": sessions, "flags": flags}


def strip_prefix(label, parent_label):
    """'QA tile tapping' under 'QA functional' reads as 'Tile tapping'."""
    word = parent_label.split()[0]
    rest = label[len(word):].strip() if label.lower().startswith(word.lower() + " ") else label
    return rest[:1].upper() + rest[1:]


def read_unit(name, ws, wsv):
    pairs = session_pairs(ws)
    groups, totals = {}, []
    for r in range(2, ws.max_row + 1):
        label = str(ws.cell(r, 1).value or "").strip()
        if not label:
            continue
        if label.lower().startswith("total"):
            totals.append((label, label[5:].strip().lower(), wsv.cell(r, 2).value))
            continue
        key = fill_key(ws.cell(r, 1))
        if key not in PEOPLE:
            raise SystemExit(f"{name}!A{r} ({label!r}) has fill {key!r}, which is not a "
                             f"known person. Add it to PEOPLE in {Path(__file__).name}.")
        groups.setdefault(key, []).append(dict(read_row(ws, r, pairs), sheet_label=label))

    activities = []
    for key, rows in groups.items():
        slot, person = PEOPLE[key]
        parent = next((x for x in rows if "(total)" in x["sheet_label"].lower()), rows[0])
        subs = [x for x in rows if x is not parent]
        label = re.sub(r"\s*\(total\)\s*$", "", parent["sheet_label"], flags=re.I)
        if slot == "cap":
            label = f"Capture ({label})"
        for s in subs:
            s["label"] = strip_prefix(s["sheet_label"], label)
            if s["sessions"] and parent["sessions"] and not all(
                    any(ps <= a and b <= pe for ps, pe in parent["sessions"])
                    for a, b in s["sessions"]):
                s["flags"].append(f"Logged {plain(fmt_windows(s['sessions']))}, outside "
                                  f"{label}'s {plain(fmt_windows(parent['sessions']))}.")
                s["outside"] = True

        other = None
        if subs:
            other = {"label": "Other functional", "state": parent["state"], "minutes": None,
                     "flags": [], "within": [s["label"] for s in subs if s["state"] == "separate"]}
            if parent["state"] == "recorded":
                known = sum(s["minutes"] for s in subs if s["state"] == "recorded")
                other["minutes"] = parent["minutes"] - known
                if other["minutes"] < 0:
                    other["flags"].append(f"The sub-tasks add up to {fmt_dur(known)}, more than "
                                          f"the {fmt_dur(parent['minutes'])} recorded for {label}.")
        activities.append(dict(parent, label=label, slot=slot, person=person,
                               subs=subs, other=other))
    for a in activities:
        if a["state"] == "pending" and a["slot"] in UNTIMED:
            a.update(state="recorded", minutes=UNTIMED[a["slot"]], untimed=True)
    activities.sort(key=lambda a: SLOTS.index(a["slot"]))
    return {"name": name, "slug": re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-"),
            "activities": activities, "sheet_totals": totals}


def check_totals(unit):
    """The sheet's own totals block, against the rows above it. A mismatch
    lands on the tile covering the same people."""
    flags = {}
    for label, word, stated in unit["sheet_totals"]:
        acts = [a for a in unit["activities"]
                if not word or a["sheet_label"].lower().startswith(word)]
        slots = frozenset(a["slot"] for a in acts) if word else frozenset(SLOTS)
        computed = sum(a["minutes"] or 0 for a in acts if not a.get("untimed"))
        if isinstance(stated, (int, float)) and round(stated) != computed:
            flags[slots] = (f"The sheet's {label!r} says {fmt_dur(round(stated))}; "
                            f"its rows add up to {fmt_dur(computed)}.")
    return flags


def mark_repeats(units):
    """A sub-task logged outside its activity is often a row carried over from
    another tab. Say so where the times match exactly."""
    for u in units:
        for a in u["activities"]:
            for s in a["subs"]:
                if not s.get("outside"):
                    continue
                same = [v["name"] for v in units if v is not u
                        for b in v["activities"] for t in b["subs"]
                        if t["sheet_label"] == s["sheet_label"] and t["sessions"] == s["sessions"]]
                if same:
                    s["flags"].append(f"{', '.join(same)} logs the same times for {s['label']}.")


def read_details(names):
    """{unit: (floor area m², bedrooms)} from UNITS_PATH. A unit it does not
    list simply has no details; a name it lists that is not a unit stops the
    build, so a typo cannot quietly drop a unit's area."""
    if not UNITS_PATH.exists():
        return {}
    out = {}
    for row in csv.DictReader(UNITS_PATH.open(newline="", encoding="utf-8-sig")):
        name = row["unit"].strip()
        if name not in names:
            raise SystemExit(f"{UNITS_PATH.name} lists {name!r}, which is not a unit in the "
                             f"workbook ({', '.join(sorted(names))}).")
        area, beds = row.get("floor_area_sqm", "").strip(), row.get("bedrooms", "").strip()
        out[name] = (float(area) if area else None, int(beds) if beds else None)
    return out


def read_units():
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("build_time_motion.py needs openpyxl: pip install openpyxl")
    wb = openpyxl.load_workbook(XLSX_PATH)
    wbv = openpyxl.load_workbook(XLSX_PATH, data_only=True)  # cached totals
    units = [read_unit(ws.title, ws, wbv[ws.title]) for ws in wb.worksheets
             if ws.title not in EXCLUDE and ws.sheet_state == "visible"]
    mark_repeats(units)
    details = read_details({u["name"] for u in units})
    for u in units:
        u["area"], u["bedrooms"] = details.get(u["name"], (None, None))
        u["total_flags"] = check_totals(u)  # against the tab as logged
        u["site"] = next((key for key, prefix, _ in SITES if u["name"].startswith(prefix)), None)
        if u["site"] is None:
            raise SystemExit(f"Workbook tab {u['name']!r} matches no project in SITES "
                             f"({', '.join(prefix for _, prefix, _ in SITES)}). "
                             f"Add one, or list the tab in EXCLUDE.")
    by_name = {u["name"]: u for u in units}
    for unit, tab in CAPTURE_FROM.items():
        if unit not in by_name or tab not in by_name:
            raise SystemExit(f"CAPTURE_FROM names {unit!r} <- {tab!r}, but the workbook has "
                             f"no such unit.")
        cap = next(a for a in by_name[tab]["activities"] if a["slot"] == "cap")
        dst = by_name[unit]
        dst["activities"] = sorted(
            [a for a in dst["activities"] if a["slot"] != "cap"] + [dict(cap, logged_on=tab)],
            key=lambda a: SLOTS.index(a["slot"]))
    for t in TESTS:
        unknown = (t.get("units") or set(t.get("logged", {}))) - set(by_name)
        if unknown:
            raise SystemExit(f"TESTS lists {t['name']!r} for {', '.join(sorted(unknown))}, "
                             f"not a unit in the workbook.")
        for unit in t.get("logged", {}):
            if test_time(t, by_name[unit])[0] <= 0:
                raise SystemExit(f"TESTS: {t['name']!r} on {unit} ends before it starts.")
    for u in units:
        for t in TESTS:
            label, note, slot = t["name"], t["note"], t["part_of"]
            a = next((a for a in u["activities"] if a["slot"] == slot), None) if slot else None
            if a is None or not applies(t, u):
                continue
            minutes, when = test_time(t, u)
            if a["other"] is None:  # no sub-tasks yet: what was logged becomes the rest
                a["other"] = {"label": "Other functional", "state": a["state"],
                              "minutes": a["minutes"], "flags": [], "within": []}
            test = {"label": label, "sheet_label": label, "state": "recorded",
                    "minutes": minutes, "sessions": [], "flags": []}
            if when:  # drawn with the activity it is part of
                test["sessions"] = [when]
                a["sessions"] = sorted(a["sessions"] + [when])
            else:
                test.update(untimed=True, note=f"{note} {UNLOGGED}.")
            a["subs"] = a["subs"] + [test]
            a.update(state="recorded", minutes=(a["minutes"] or 0) + minutes,
                     includes=a.get("includes", []) + [(label, minutes, when is None)])
    missing = HIDDEN - set(by_name)
    if missing:
        raise SystemExit(f"HIDDEN names {', '.join(sorted(missing))}, not a unit in the workbook.")
    units = [u for u in units if u["name"] not in HIDDEN]
    # T240 before T241, wherever the tabs sit in the workbook
    units.sort(key=lambda u: [int(t) if t.isdigit() else t.lower()
                              for t in re.split(r"(\d+)", u["name"])])
    return units


# ----------------------------------------------------------------- render ---

def warn(flags):
    if not flags:
        return ""
    tip = " ".join(flags)
    return (f'<span class="warn" tabindex="0" role="note" aria-label="{esc(tip)}" '
            f'data-tip="{esc(tip)}">!</span>')


def pill(state):
    return f'<span class="st st-{state}">{STATE_TEXT[state]}</span>'


def short_label(a):
    """'MEP specialist (above ceiling)' as 'MEP specialist'; capture as its platform."""
    if a["slot"] == "cap":
        return a["label"].replace("Capture (", "").rstrip(")")
    return re.sub(r"\s*\(.*\)$", "", a["label"])


def minutes_of(acts, slots):
    return sum(a["minutes"] or 0 for a in acts if a["slot"] in slots)


def unit_total(u):
    """Everyone's minutes, tests that are part of someone's activity included."""
    return minutes_of(u["activities"], SLOTS)


def per_area(mins, area):
    """Minutes per PER_AREA m² of floor area, or None without an area."""
    return round(mins * PER_AREA / area) if area else None


def fmt_area(a):
    return f"{a:,.0f} m&sup2;" if a == int(a) else f"{a:,.1f} m&sup2;"


PER_LABEL = f"per {PER_AREA} m&sup2;"
# for the uppercase tile labels, where m² must not become M²
PER_KEY = f'Per {PER_AREA} <span class="u">m&sup2;</span>'


def unit_facts(u, short=False):
    """'290 m² · 6 bedrooms', or as much of it as is known."""
    bits = []
    if u["area"]:
        bits.append(fmt_area(u["area"]))
    if u["bedrooms"]:
        bits.append(f'{u["bedrooms"]} bed' if short else
                    f'{u["bedrooms"]} bedroom{"s" if u["bedrooms"] != 1 else ""}')
    return " &middot; ".join(bits)


def logged_times(unit):
    return [t for a in unit["activities"] for x in [a] + a["subs"]
            for s in x["sessions"] for t in s]


def axis_for(unit):
    times = logged_times(unit) or [9 * 60, 17 * 60]  # nothing timed yet: a working day
    lo = (min(times) // 60) * 60
    hi = max(lo + 120, math.ceil(max(times) / 60) * 60)
    step = 30 if hi - lo <= 180 else 60
    return lo, hi, list(range(lo, hi + 1, step))


def pos(t, axis):
    lo, hi, _ = axis
    return 100 * (t - lo) / (hi - lo)


def track(axis, bars=(), wash=(), note=""):
    ticks = "".join(f'<i class="tk" style="left:{pos(t, axis):.3f}%"></i>' for t in axis[2])
    spans = "".join(
        f'<span class="wash" style="left:{pos(s, axis):.3f}%;width:{pos(e, axis) - pos(s, axis):.3f}%"></span>'
        for s, e in wash)
    spans += "".join(
        f'<span class="bar" style="left:{pos(s, axis):.3f}%;width:{pos(e, axis) - pos(s, axis):.3f}%"'
        f' title="{esc(tip)}"></span>' for s, e, tip in bars)
    if note:  # words, not a timeline: no gridlines behind them
        return f'<div class="c-trk"><div class="trk"><span class="trk-note">{note}</span></div></div>'
    return f'<div class="c-trk"><div class="trk">{ticks}{spans}</div></div>'


def time_cell(x):
    if x["state"] == "recorded":
        return f'<div class="c-time"><span class="mins mono">{fmt_dur(x["minutes"])}</span>{warn(x["flags"])}</div>'
    if x["state"] == "text":
        return f'<div class="c-time"><span class="mins">{esc(x["text"])}</span></div>'
    return f'<div class="c-time">{pill(x["state"])}{warn(x["flags"])}</div>'


NONE = "&mdash;"  # no clock times

CHEVRON = ('<svg class="chev" viewBox="0 0 16 16" aria-hidden="true">'
           '<path d="M6 4l4 4-4 4" fill="none" stroke="currentColor" stroke-width="1.8" '
           'stroke-linecap="round" stroke-linejoin="round"/></svg>')


def render_activity(unit, a, axis):
    bars = [(s, e, f'{unit["name"]} · {a["label"]} · {plain(fmt_windows([(s, e)]))}')
            for s, e in a["sessions"]]
    logged = fmt_windows(a["sessions"], "<br>") if a["sessions"] else NONE
    hint = ""
    if a["subs"]:
        n = len(a["subs"])
        hint = (f'<span class="hint">{n} sub-task{"s" if n > 1 else ""} '
                f'+ other functional</span>')
    name = (f'<div class="c-name">{CHEVRON if a["subs"] else "<i class=chev-sp></i>"}'
            f'<i class="sw"></i><span class="nm"><b>{esc(a["label"])}</b>'
            f'<span class="who">{esc(a["person"])}</span>{hint}</span></div>')
    row = (f'{name}<div class="c-log mono{" none" if logged == NONE else ""}">{logged}</div>{time_cell(a)}'
           f'{track(axis, bars, note=f"{UNLOGGED}." if a.get("untimed") else "")}')
    if not a["subs"]:
        return f'<div class="act p-{a["slot"]}"><div class="row">{row}</div></div>'

    subs = []
    for s in a["subs"]:
        bars = [(x, y, f'{unit["name"]} · {s["label"]} · {plain(fmt_windows([(x, y)]))}')
                for x, y in s["sessions"]]
        subs.append(
            f'<div class="row sub"><div class="c-name"><span class="nm">{esc(s["label"])}</span></div>'
            + (f'<div class="c-log mono">{fmt_windows(s["sessions"], "<br>")}</div>' if s["sessions"]
               else f'<div class="c-log none">{NONE}</div>')
            + f'{time_cell(s)}'
            + (track(axis, note=esc(s["note"])) if s.get("note") else track(axis, bars, a["sessions"]))
            + '</div>')
    o = a["other"]
    within = (f'<span class="who">Includes {esc(", ".join(o["within"]).lower())}, '
              f'not recorded separately</span>' if o["within"] else "")
    subs.append(
        f'<div class="row sub other"><div class="c-name"><span class="nm">Other functional'
        f'<span class="who">{esc(a["label"])} time not covered by the sub-tasks</span>{within}</span></div>'
        f'<div class="c-log">Rest of session</div>{time_cell(o)}'
        f'{track(axis, (), a["sessions"])}</div>')
    return (f'<details class="act p-{a["slot"]}"><summary class="row">{row}</summary>'
            f'<div class="subs">{"".join(subs)}</div></details>')


def render_tests(u, axis):
    out = []
    for t in TESTS:
        if not applies(t, u):
            continue
        label, who, status, note, slot = t["name"], t["who"], t["status"], t["note"], t["part_of"]
        minutes, when = test_time(t, u)
        part_of = next((a["label"] for a in u["activities"] if a["slot"] == slot), None)
        if part_of:
            note = f"{note} Counted in {part_of} time."
        if not when:
            note = f"{note} {UNLOGGED}."
        logged = (f'<div class="c-log mono">{fmt_windows([when])}</div>' if when
                  else f'<div class="c-log none">{NONE}</div>')
        out.append(
            f'<div class="act test{f" part-of p-{slot}" if part_of else ""}"><div class="row">'
            f'<div class="c-name"><i class="chev-sp"></i><i class="sw sw-test"></i>'
            f'<span class="nm"><b>{esc(label)}</b>'
            + (f'<span class="who">{esc(who)}</span>' if who else "") + '</span></div>'
            + logged +
            f'<div class="c-time stack"><span class="mins mono">{fmt_dur(minutes)}</span>'
            f'<span class="st st-done">{esc(status)}</span></div>'
            f'{track(axis, note=esc(note))}</div></div>')
    return "".join(out)


TILES = [
    ("Recorded time", frozenset(SLOTS)),
    ("MEP", frozenset({"mepg", "meps"})),
    ("QA", frozenset({"qa"})),
    ("Capture", frozenset({"cap"})),
]


def tile_sub(acts, slots):
    """What a tile is made of, or why it is short."""
    inside = [a for a in acts if a["slot"] in slots]
    if len(slots) == len(SLOTS):
        n = len({a["slot"] for a in inside if a["state"] == "recorded"})
        return f"{n} {'person' if n == 1 else 'people'}, all activities"
    bits = []
    for a in inside:
        short = esc(short_label(a))
        incl = and_join(t.lower() for t, *_ in a.get("includes", []))
        if a["state"] == "recorded" and len(inside) > 1:
            bits.append(f'{short} {fmt_dur(a["minutes"])}' + (f" incl. {incl}" if incl else ""))
        elif a["state"] == "recorded":
            bits.append(short + (f", incl. {incl}" if incl else ""))
        else:
            bits.append(f'{short} {STATE_TEXT.get(a["state"], "").lower()}')
    return " + ".join(bits) or "&mdash;"


def render_unit(u):
    acts, axis = u["activities"], axis_for(u)
    tiles = []
    for label, slots in TILES:
        inside = [a for a in acts if a["slot"] in slots]
        if inside and all(a["state"] != "recorded" for a in inside):
            v = "&mdash;"
        elif len(slots) == len(SLOTS):
            v = fmt_dur(unit_total(u))
        else:
            v = fmt_dur(minutes_of(acts, slots))
        tiles.append(f'<div class="tile"><span class="k">{label}{warn([u["total_flags"][slots]]) if slots in u["total_flags"] else ""}</span>'
                     f'<span class="v mono">{v}</span><span class="s">{tile_sub(acts, slots)}</span></div>')
    if u["area"]:
        tiles.append(f'<div class="tile"><span class="k">{PER_KEY}</span>'
                     f'<span class="v mono">{fmt_dur(per_area(unit_total(u), u["area"]))}</span>'
                     f'<span class="s">recorded time over {fmt_area(u["area"])}</span></div>')
    facts = unit_facts(u)
    facts = f'<span class="ufacts">{facts}</span>' if facts else ""

    times = logged_times(u)
    cap = next((a for a in acts if a["slot"] == "cap"), None)
    bits = []
    if cap:
        bits.append(f"capture on <b>{esc(short_label(cap))}</b>"
                    + (" &middot; pending" if cap["state"] == "pending" else ""))
    if times:
        bits.append(f"first to last logged activity <b>{fmt_clock(min(times))}&ndash;"
                    f"{fmt_clock(max(times))}</b>")
    else:
        bits.append("no activity logged yet")
    sub = " &middot; ".join(bits)
    sub = sub[:1].upper() + sub[1:]

    lo, hi, ticks = axis
    every = math.ceil(len(ticks) / 5)  # a gridline every step, but at most ~5 labels
    labels = "".join(
        f'<span style="left:{pos(t, axis):.3f}%"'
        + (' class="first"' if t == lo else ' class="last"' if t == hi else "")
        + f'>{fmt_clock(t)}</span>' for t in ticks[::every])
    return f"""
<article class="card unit-card" id="u-{esc(u["slug"])}" data-unit="{esc(u["slug"])}">
  <div class="uhead">
    <div class="uname">
      <h3>{esc(u["name"])}</h3>{facts}
      <p class="usub">{sub}</p>
    </div>
    <div class="tiles" style="--n:{len(tiles)}">{"".join(tiles)}</div>
  </div>
  <div class="acts">
    <div class="row head" aria-hidden="true">
      <div class="c-name">Activity</div><div class="c-log">Logged</div>
      <div class="c-time">Time</div>
      <div class="c-trk"><div class="axis">{labels}</div></div>
    </div>
    {"".join(render_activity(u, a, axis) for a in acts)}
    <div class="tests-head">Tests</div>
    {render_tests(u, axis)}
  </div>
</article>"""


def render_compare(units):
    """Recorded time per unit, split by person, either as it stands or per
    PER_AREA m² of floor area. Either way bar length is comparable across
    units; the card below each one breaks it down."""
    sized = [u for u in units if u["area"]]
    totals = {u["slug"]: unit_total(u) for u in units}
    norms = {u["slug"]: per_area(totals[u["slug"]], u["area"]) for u in units}
    peak = max(totals.values()) or 1
    npeak = max((n for n in norms.values() if n), default=0) or 1
    names = dict(PEOPLE.values())
    labels = {}
    for u in units:
        for a in u["activities"]:
            labels.setdefault(a["slot"], a["label"] if a["slot"] != "cap" else "Capture")

    def both(mins, area):
        """The figure in each mode; only one shows at a time."""
        n = per_area(mins, area)
        return (f'<span class="v-abs">{fmt_dur(mins)}</span>'
                f'<span class="v-area">{fmt_dur(n) if n is not None else "&mdash;"}</span>')

    rows = []
    for u in units:
        by = {a["slot"]: a for a in u["activities"]}

        def tip(what, mins, extra=""):
            n = per_area(mins, u["area"])
            return esc(f'{u["name"]} · {what} · {fmt_dur(mins)}'
                       + (f" ({fmt_dur(n)} per {PER_AREA} m\u00b2)" if n is not None else "")
                       + extra)
        def extra(a):
            if a.get("untimed"):
                return " · " + UNLOGGED.lower()
            return "".join(f" · includes {t.lower()} {fmt_dur(m)}"
                           + (f" ({UNLOGGED.lower()})" if untimed else "")
                           for t, m, untimed in a.get("includes", []))
        segs = "".join(
            f'<i class="p-{s}" style="flex:{by[s]["minutes"]}" title="'
            + tip(f"{labels[s]} ({names[s]})", by[s]["minutes"], extra(by[s]))
            + '"></i>'
            for s in SLOTS if s in by and by[s]["state"] == "recorded")
        notes = []
        for state, word in (("pending", "pending"), ("separate", "not recorded separately")):
            missing = [("capture" if s == "cap" else short_label(by[s]))
                       for s in SLOTS if s in by and by[s]["state"] == state]
            if missing:
                notes.append(f'{" and ".join(missing)} {word}')
        note = esc("; ".join(notes))
        if sized and not u["area"]:
            note = (f'<span class="v-abs">{note}</span>'
                    f'<span class="v-area">no floor area given</span>')
        note = f'<span class="cmp-note">{note}</span>' if note else ""
        n = norms[u["slug"]]
        facts = unit_facts(u, short=True)
        rows.append(f"""
    <div class="cmp-row" data-unit="{esc(u["slug"])}">
      <a class="cmp-u" href="#u-{esc(u["slug"])}">{esc(u["name"])}"""
                    + (f'<span class="cmp-facts">{facts}</span>' if facts else "") + f"""</a>
      <div class="cmp-track">
        <div class="cmp-bar" style="--fa:{totals[u["slug"]] / peak:.4f};--fn:{(n or 0) / npeak:.4f}">{segs}</div>
        <span class="cmp-lab mono">{both(totals[u["slug"]], u["area"])}</span>{note}
      </div>
    </div>""")

    # per person across units; per area over the units that have one
    area_sum = sum(u["area"] for u in sized)
    def legend_value(mins_all, mins_sized):
        n = round(mins_sized * PER_AREA / area_sum) if area_sum else None
        return (f'<span class="v-abs">{fmt_dur(mins_all)}</span>'
                f'<span class="v-area">{fmt_dur(n) if n is not None else "&mdash;"}</span>')
    legend = "".join(
        f'<span class="p-{s}"><i class="sw"></i>{esc(labels[s])} '
        + ("" if labels[s].startswith(names[s]) else f'<span class="lg-who">{esc(names[s])}</span> ')
        + '<b class="mono">' + legend_value(sum(minutes_of(u["activities"], {s}) for u in units),
                                            sum(minutes_of(u["activities"], {s}) for u in sized))
        + '</b></span>'
        for s in SLOTS if s in labels)

    folded = "".join(
        f' {esc(labels[t["part_of"]])} includes the {esc(t["name"].lower())}'
        + ("" if len(on) == len(units) else f' on {esc(and_join(short_unit(u["name"]) for u in on))}')
        + "."
        for t in TESTS if t["part_of"] in labels
        for on in [[u for u in units if applies(t, u)]] if on)
    toggle = sub_area = ""
    if sized:
        toggle = f"""
  <div class="modes" role="group" aria-label="Show recorded time as">
    <button type="button" class="mode" data-mode="abs" aria-pressed="true">Total time</button>
    <button type="button" class="mode" data-mode="area" aria-pressed="false">{PER_LABEL.capitalize()}</button>
  </div>"""
        sub_area = (f'<span class="v-area">The same minutes divided by each unit&rsquo;s floor area, as '
                    f'time {PER_LABEL}, so units of different sizes compare. '
                    + ("" if len(sized) == len(units) else "Units without a floor area have no bar. ")
                    + "Unit names jump to the breakdown below.</span>")
    return f"""
<section class="card cmp-card" aria-labelledby="cmp-h">
  <div class="card-head cmp-head">
    <div>
      <h2 id="cmp-h">Recorded time by unit</h2>
      <p class="card-sub"><span class="v-abs">Minutes logged on site, split by the person who logged
      them.{folded} Bar length is comparable across units. Unit names jump to the breakdown
      below.</span>{sub_area}</p>
    </div>{toggle}
  </div>
  <div class="legend">{legend}</div>
  <div class="cmp">{"".join(rows)}</div>
</section>"""


def page_file(key):
    return next(f for slug, _, f in theme.PAGES if slug == key)


def redirect(units, everyone):
    """index.html used to be the activity log, so old links (…/#BL515) arrive
    here. Send any hash this page does not know on to wherever it lives."""
    if page_file(units[0]["site"]) != "index.html":
        return ""
    known = ",".join(f'"{u["slug"]}"' for u in units)
    elsewhere = ",".join(f'"{u["slug"]}":"{page_file(u["site"])}"'
                         for u in everyone if u["site"] != units[0]["site"])
    return f"""<script>
/* Links into the archived activity log used to land here (…/#BL515). Send them on. */
(function(){{
  var h = decodeURIComponent(location.hash.slice(1)).replace(/^u-/, '');
  var elsewhere = {{{elsewhere}}};
  if (h && [{known}].indexOf(h) < 0)
    location.replace((elsewhere[h] || '{page_file("activity")}') + location.hash);
}})();
</script>"""


def render(site, units, everyone):
    key, prefix, heading = site
    many = len(units) > 1
    all_acts = [a for u in units for a in u["activities"]]
    total = sum(unit_total(u) for u in units)
    on_page = [t for t in TESTS if any(applies(t, u) for u in units)]

    def included(slots):
        """Tests on this page that are part of these people's activities."""
        return [t["name"].lower() for t in on_page if t["part_of"] in slots]
    sized = [u for u in units if u["area"]]
    area = sum(u["area"] for u in sized)
    caps = [u for u in units if any(a["slot"] == "cap" and a["state"] == "recorded"
                                    for a in u["activities"])]
    strip = "".join(f'<div><span class="k">{k}</span><span class="v mono">{v}</span>'
                    f'<span class="s">{s}</span></div>' for k, v, s in [
        ("Units", str(len(units)),
         ", ".join(esc(re.sub(rf"^{re.escape(prefix)}[\s-]*", "", u["name"])) for u in units)),
        ("Recorded time", fmt_dur(total), "person-time on site, all units"),
        ("MEP", fmt_dur(minutes_of(all_acts, {"mepg", "meps"})),
         "general and specialist" + (f", {and_join(included({'mepg', 'meps'}))} included"
                                     if included({"mepg", "meps"}) else "")),
        ("QA", fmt_dur(minutes_of(all_acts, {"qa"})),
         "functional, sub-tasks" + "".join(f" and {n}" for n in included({"qa"})) + " included"),
        ("Capture", fmt_dur(minutes_of(all_acts, {"cap"})),
         f"{len(caps)} of {len(units)} units captured"),
    ] + ([(PER_KEY,
           fmt_dur(round(sum(unit_total(u) for u in sized) * PER_AREA / area)),
           f"recorded time over {fmt_area(area)}"
           + ("" if len(sized) == len(units) else f", the {len(sized)} units with a floor area"))]
         if sized else []))

    chips = "".join(
        f'<button class="chip" role="tab" aria-selected="false" data-f="{esc(u["slug"])}">'
        f'{esc(u["name"])}</button>' for u in units)
    filters = f"""
<div class="filters" role="tablist" aria-label="Filter by unit">
  <span class="flabel">Unit</span>
  <div class="chips">
    <button class="chip" role="tab" aria-selected="true" data-f="all">All</button>
    {chips}
  </div>
  <span class="showing" id="showing">Showing all {len(units)} units</span>
  <button class="expand" type="button" id="expand" aria-pressed="false">Open all sub-tasks</button>
</div>""" if many else """
<div class="filters">
  <button class="expand" type="button" id="expand" aria-pressed="false">Open all sub-tasks</button>
</div>"""
    # one unit: its card's tiles are the summary, and a one-bar chart says nothing
    tiles = strip.count('<span class="k">')
    summary = (f'<div class="strip{" six" if tiles == 6 else ""}">{strip}</div>\n\n'
               f'{render_compare(units)}' if many else "")
    per_note = (f"""<li><b>{PER_LABEL.capitalize()}</b> is a unit&rsquo;s recorded time divided by its
      floor area, so units of different sizes compare. Switch the chart
      between total time and {PER_LABEL} with the buttons above it. Floor areas and bedrooms come
      from <span class="mono">data/{UNITS_PATH.name}</span>.</li>
    """ if sized else "")
    labels = {a["slot"]: a["label"] for u in units for a in u["activities"]}
    def test_sentence(t):
        on = [u for u in units if applies(t, u)]
        times = [test_time(t, u) for u in on]
        if "logged" in t:
            scope = (and_join(f'{short_unit(u["name"])} {fmt_dur(m)}' for u, (m, _) in zip(on, times))
                     + ", from its logged start and end")
        else:
            scope = (fmt_dur(t["minutes"])
                     + ("" if len(on) == len(units) else ", " + and_join(short_unit(u["name"]) for u in on))
                     + ", " + UNLOGGED.lower())
        out = f'The <b>{esc(t["name"].lower())}</b> ({esc(scope)}): {esc(t["note"][:1].lower() + t["note"][1:])} '
        out += (f'Its time is part of {esc(labels[t["part_of"]])} &mdash; a sub-task in its dropdown '
                f'&mdash; and of the unit&rsquo;s recorded time.' if t["part_of"] in labels
                else "It is not part of the unit totals.")
        return out
    tests_note = " ".join(test_sentence(t) for t in on_page)
    skipped = sorted(t for t in EXCLUDE if t.startswith(prefix))
    skipped = (f"""
    <li>The workbook's {", ".join(f'<span class="mono">{esc(t)}</span>' for t in skipped)}
      {"tab is an earlier take" if len(skipped) == 1 else "tabs are earlier takes"} and
      {"is" if len(skipped) == 1 else "are"} not shown.</li>""" if skipped else "")

    masthead = theme.masthead(
        key, f"Pilot time and motion: {esc(heading)}",
        f"How long each activity took on each {esc(heading)} pilot unit, timed on site. "
        "One colour per person; open an activity for its sub-tasks.",
        f"Source <b>QA/QC time + motion</b><br>"
        f"Units <b>{len(units)}</b> &middot; Generated <b>{GENERATED:%d %b %Y}</b><br>"
        f"Times as recorded onsite")

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex">
<title>DAMAC snagging pilot &mdash; time and motion, {esc(heading)}</title>
<meta name="description" content="Time and motion for each {esc(heading)} unit in the DAMAC x CONXAI snagging pilot.">
{redirect(units, everyone)}
<style>
{theme.CHROME_CSS}

/* ------------------------------------------ the time and motion page --- */
:root{{
  /* one per person, the sheet's hues; validated as a categorical set */
  --p-mepg:#cf7b26; --p-meps:#065da0; --p-qa:#a672cb; --p-cap:#1e7546;
}}
.p-mepg{{--pc:var(--p-mepg)}} .p-meps{{--pc:var(--p-meps)}}
.p-qa{{--pc:var(--p-qa)}} .p-cap{{--pc:var(--p-cap)}}
.sw{{display:block;width:10px;height:10px;border-radius:3px;background:var(--pc);flex:none}}

.legend{{display:flex;flex-wrap:wrap;gap:8px 20px;font-size:12.4px;color:var(--ink2);margin-bottom:16px}}
.legend>span{{display:flex;align-items:center;gap:7px}}
.legend .lg-who{{color:var(--muted)}}
.legend b{{color:var(--ink);font-weight:650}}
.cmp{{display:flex;flex-direction:column;gap:8px}}
.cmp-row{{display:grid;grid-template-columns:104px 1fr;gap:14px;align-items:center;transition:opacity .15s}}
.cmp-row.dim{{opacity:.35}}
.cmp-u{{font-size:12.8px;font-weight:650;color:var(--ink);text-decoration:none;display:flex;flex-direction:column}}
.k .u{{text-transform:none}}
/* six summary tiles: one row when wide, else even rows of three or two */
.strip.six{{grid-template-columns:repeat(6,1fr)}}
@media (max-width:1100px){{.strip.six{{grid-template-columns:repeat(3,1fr)}}}}
@media (max-width:560px){{.strip.six{{grid-template-columns:repeat(2,1fr)}}}}
.cmp-facts{{font-size:11px;font-weight:500;color:var(--muted);line-height:1.3}}
.cmp-head{{display:flex;flex-wrap:wrap;gap:12px 24px;align-items:flex-start;justify-content:space-between}}
.cmp-head>div:first-child{{flex:1 1 420px}}
.modes{{display:flex;border:1px solid var(--line);border-radius:8px;overflow:hidden;background:var(--surface);flex:none}}
.mode{{font:inherit;font-size:12.5px;padding:7px 12px;border:none;background:none;cursor:pointer;color:var(--ink2)}}
.mode+.mode{{border-left:1px solid var(--line)}}
.mode:hover{{background:var(--panel);color:var(--ink)}}
.mode[aria-pressed="true"]{{background:var(--ink);color:#fff;font-weight:600}}
.mode:focus-visible{{outline:2px solid #1b6fc4;outline-offset:-2px}}
.cmp-card .v-area,.cmp-card.area .v-abs{{display:none}}
.cmp-card.area .v-area{{display:inline}}
.cmp-u:hover{{text-decoration:underline}}
.cmp-track{{display:flex;align-items:center;gap:10px;min-width:0}}
.cmp-bar{{display:flex;gap:2px;height:22px;flex:none;--f:var(--fa);width:calc((100% - 250px) * var(--f));
  transition:width .25s ease}}
.cmp-card.area .cmp-bar{{--f:var(--fn)}}
@media (prefers-reduced-motion:reduce){{.cmp-bar{{transition:none}}}}
.cmp-bar i{{display:block;background:var(--pc);min-width:2px}}
.cmp-bar i:last-child{{border-radius:0 4px 4px 0}}
.cmp-lab{{font-size:12.8px;font-weight:700;white-space:nowrap}}
.cmp-note{{font-size:11.6px;color:var(--muted);line-height:1.35;min-width:0;flex:1 1 0}}

.filters{{display:flex;flex-wrap:wrap;gap:9px;align-items:center;margin-bottom:20px;
  position:sticky;top:0;z-index:5;background:var(--page);padding:10px 0 11px}}
.filters::after{{content:"";position:absolute;left:0;right:0;bottom:0;height:1px;background:var(--line)}}
.flabel{{font-size:11px;letter-spacing:.07em;text-transform:uppercase;color:var(--muted);font-weight:650}}
.chips{{display:flex;flex-wrap:wrap;gap:0;border:1px solid var(--line);border-radius:8px;
  overflow:hidden;background:var(--surface)}}
.chip,.expand{{font:inherit;font-size:12.5px;padding:8px 13px;border:none;background:none;cursor:pointer;
  color:var(--ink2)}}
.chip{{border-right:1px solid var(--line)}}
.chip:last-child{{border-right:none}}
.chip:hover,.expand:hover{{background:var(--panel);color:var(--ink)}}
.chip[aria-selected="true"]{{background:var(--ink);color:#fff;font-weight:600}}
.chip:focus-visible,.expand:focus-visible{{outline:2px solid #1b6fc4;outline-offset:-2px}}
.expand{{margin-left:auto;border:1px solid var(--line);border-radius:8px;background:var(--surface)}}
.showing{{font-size:12.5px;color:var(--muted)}}

.uhead{{display:flex;flex-wrap:wrap;gap:18px;align-items:flex-start;justify-content:space-between;
  padding-bottom:16px;border-bottom:1px solid var(--line);margin-bottom:6px}}
.uname h3{{font-size:26px;margin:0;font-weight:700;letter-spacing:-.02em;display:inline-block;vertical-align:baseline}}
.ufacts{{display:inline-block;margin-left:10px;vertical-align:4px;font-size:12px;font-weight:600;color:var(--ink2);
  background:var(--panel);border:1px solid var(--line2);border-radius:999px;padding:2px 9px;white-space:nowrap}}
.usub{{color:var(--ink2);font-size:12.8px;margin:5px 0 0}}
.usub b{{color:var(--ink);font-weight:600}}
.tiles{{display:grid;grid-template-columns:repeat(var(--n,4),minmax(104px,1fr));gap:9px;flex:0 1 calc(var(--n,4) * 150px)}}
.tile{{background:var(--panel);border:1px solid var(--line2);border-radius:9px;padding:10px 12px}}
.tile .k{{display:block;font-size:10.5px;text-transform:uppercase;letter-spacing:.07em;
  color:var(--muted);font-weight:650;line-height:1.35}}
.tile .v{{display:block;font-size:19px;font-weight:700;letter-spacing:-.02em;margin-top:3px}}
.tile .s{{display:block;font-size:11.2px;color:var(--ink2);margin-top:3px;line-height:1.45}}

/* activity rows: name | logged | time | timeline */
.row{{display:grid;grid-template-columns:minmax(210px,1.3fr) 96px 158px minmax(220px,2.3fr);
  gap:0 16px;align-items:stretch}}
.row>div{{display:flex;align-items:center;min-width:0}}
.row>div:not(.c-trk){{padding:9px 0}}
.act{{border-top:1px solid var(--line2)}}
.row.head{{font-size:10.5px;text-transform:uppercase;letter-spacing:.07em;color:var(--muted);
  font-weight:650}}
.row.head>div:not(.c-trk){{padding:12px 0 6px}}
.c-name{{gap:8px}}
.nm{{display:flex;flex-direction:column;font-size:13px;line-height:1.35}}
.nm b{{font-weight:650}}
.who,.hint{{font-size:11.5px;color:var(--muted)}}
.c-log{{font-size:12px;color:var(--ink2);line-height:1.5}}
.c-time{{justify-content:flex-end;text-align:right;gap:0}}
.row>.c-time.stack{{flex-direction:column;align-items:flex-end;justify-content:center;gap:4px}}
.mins{{font-size:13px;font-weight:650}}
.chev{{width:16px;height:16px;flex:none;color:var(--muted);transition:transform .15s}}
.chev-sp{{width:16px;flex:none}}
details[open]>summary .chev{{transform:rotate(90deg)}}
summary.row{{cursor:pointer;list-style:none}}
summary.row::-webkit-details-marker{{display:none}}
summary.row:hover{{background:var(--panel)}}
summary.row:hover .chev{{color:var(--ink)}}
summary.row:focus-visible{{outline:2px solid #1b6fc4;outline-offset:-2px}}
.subs{{background:var(--panel);border-top:1px solid var(--line2)}}
.row.sub>.c-name{{padding-left:24px;position:relative}}
.row.sub>.c-name::before{{content:"";position:absolute;left:7px;top:0;bottom:0;width:2px;background:var(--pc);opacity:.45}}
.row.sub .nm{{font-size:12.6px}}
.row.sub+.row.sub{{border-top:1px solid var(--line2)}}
.row.other .nm{{font-weight:600}}

.st{{display:inline-flex;align-items:center;gap:6px;font-size:11px;font-weight:600;padding:2px 9px;
  border-radius:999px;white-space:nowrap;line-height:1.55}}
.st-pending{{background:var(--surface);color:var(--ink2);border:1px solid var(--line)}}
.st-pending::before{{content:"";width:6px;height:6px;border-radius:50%;border:1.5px solid var(--muted)}}
.st-separate{{background:var(--line2);color:var(--ink2)}}
.st-done{{background:var(--ink);color:#fff}}
.st-done::before{{content:"\\2713";font-size:10px}}

.c-trk{{position:relative}}
.trk{{position:relative;width:100%;min-height:40px;height:100%}}
.tk{{position:absolute;top:0;bottom:0;width:1px;background:var(--line2)}}
.bar,.wash{{position:absolute;top:50%;height:12px;margin-top:-6px;border-radius:3px;background:var(--pc);min-width:3px}}
.wash{{opacity:.18}}
.row.sub .bar{{height:10px;margin-top:-5px}}
.trk-note{{position:absolute;left:10px;right:0;top:50%;transform:translateY(-50%);font-size:11.8px;
  color:var(--ink2);line-height:1.4}}
.axis{{position:relative;width:100%;height:100%;min-height:30px}}
.axis span{{position:absolute;bottom:6px;transform:translateX(-50%);font-family:ui-monospace,monospace;
  font-size:10.5px;letter-spacing:0;text-transform:none;font-weight:500;color:var(--muted)}}
.axis span.first{{transform:none}}
.axis span.last{{transform:translateX(-100%)}}
.tests-head{{font-size:10.5px;text-transform:uppercase;letter-spacing:.07em;color:var(--muted);
  font-weight:650;padding:16px 0 6px;border-top:1px solid var(--line)}}
.act.test .row>div:not(.c-trk){{padding:8px 0}}
.act.test .trk{{min-height:36px}}
.sw-test{{background:none;border:1.5px solid var(--muted)}}
.act.test.part-of .sw-test{{background:var(--pc);border-color:var(--pc)}}

@media (max-width:860px){{
  .row{{grid-template-columns:1fr auto;grid-template-areas:"name time" "log log" "trk trk";gap:0 12px}}
  .c-name{{grid-area:name}} .c-time{{grid-area:time}} .c-log{{grid-area:log}} .c-trk{{grid-area:trk}}
  .row>div.c-log{{padding:0 0 4px 24px}}
  .row.sub>div.c-log{{padding-left:24px}}
  .row.head .c-log,.row>div.c-log.none{{display:none}}
  .trk{{min-height:26px}}
  .trk-note{{position:static;transform:none;display:block;padding:0 0 10px 24px}}
  .act.test .trk{{min-height:0}}
  .axis span:nth-child(even){{display:none}}
}}
@media (max-width:720px){{
  .tiles{{grid-template-columns:repeat(2,minmax(110px,1fr));flex:1 1 100%}}
  .filters{{position:static}}
  .expand{{margin-left:0}}
  .cmp-row{{grid-template-columns:1fr;gap:4px}}
  .cmp-track{{flex-wrap:wrap;gap:4px 10px}}
  /* the bar takes its own line; total and note sit underneath */
  .cmp-bar{{height:18px;width:calc(100% * var(--f));margin-right:calc(100% * (1 - var(--f)))}}
}}
@media print{{
  .filters{{display:none}}
}}
</style>
</head>
<body>
<div class="wrap">

{masthead}

{summary}
{filters}

{"".join(render_unit(u) for u in units)}

<div class="notes">
  <h2>Reading the page</h2>
  <ul>
    <li><b>Each colour is a different person</b>, as in the source sheet: the MEP inspector, the MEP
      specialist, the QA inspector and whoever ran the capture. Times are the minutes each person
      logged, so a unit's recorded time is person-time, not time on site.</li>
    <li><b>MEP general</b> and <b>QA functional</b> are totals for the whole activity. Open one for its
      sub-tasks; <b>Other functional</b> is the total less the sub-tasks that were timed, i.e. the
      time they do not cover. Sub-tasks share the activity's colour, and the faint band behind them
      is the activity's own session.</li>
    <li><b>Not recorded separately</b> means the sheet logs 0 minutes or says so: the time is inside
      the activity total, and so inside Other functional. <b>Activity pending</b> means the sheet's row
      is blank.</li>
    <li>Where the sheet leaves <b>MEP specialist</b> blank, it is shown as
      {fmt_dur(UNTIMED["meps"])}, marked <i>{UNLOGGED.lower()}</i>. That time counts towards the
      unit's MEP and recorded time.</li>
    {per_note}<li><b>Tests</b> are not in the sheet. {tests_note}</li>
    <li>A <span class="warn">!</span> marks minutes that do not agree with the clock times logged beside
      them, or a sub-task logged outside its activity. Hover it for the detail. Differences of
      {ROUNDING} minutes or less are rounding and are not marked. Otherwise minutes are shown exactly
      as the sheet records them, and every total is built from them &mdash; nothing has been
      corrected.</li>{skipped}
  </ul>
</div>

<p class="foot">
  Generated from <span class="mono">data/qaqc-time-motion.xlsx</span> by
  <span class="mono">build_time_motion.py</span> &middot; {GENERATED:%d %b %Y}.
  Re-run <span class="mono">build.py</span> after updating the workbook to refresh this page.
</p>

</div>
<script>
(function(){{
  var chips = Array.prototype.slice.call(document.querySelectorAll('.chip'));
  var cards = Array.prototype.slice.call(document.querySelectorAll('.unit-card'));
  var rows  = Array.prototype.slice.call(document.querySelectorAll('.cmp-row'));
  var out   = document.getElementById('showing');
  var known = chips.map(function(c){{ return c.dataset.f; }});
  var names = {{}};
  chips.forEach(function(c){{ names[c.dataset.f] = c.textContent; }});

  function apply(f, push){{
    chips.forEach(function(c){{ c.setAttribute('aria-selected', String(c.dataset.f === f)); }});
    cards.forEach(function(c){{ c.hidden = !(f === 'all' || c.dataset.unit === f); }});
    rows.forEach(function(r){{ r.classList.toggle('dim', f !== 'all' && r.dataset.unit !== f); }});
    if (out) out.textContent = f === 'all' ? 'Showing all ' + cards.length + ' units' : 'Showing ' + names[f] + ' only';
    if (push) history.replaceState(null, '', f === 'all' ? location.pathname : '#' + f);
  }}
  function fromHash(){{
    var h = decodeURIComponent(location.hash.replace(/^#(u-)?/, ''));
    return known.indexOf(h) > -1 ? h : 'all';
  }}

  chips.forEach(function(c){{
    c.addEventListener('click', function(){{ apply(c.dataset.f, true); }});
  }});
  window.addEventListener('hashchange', function(){{ apply(fromHash(), false); }});
  apply(fromHash(), false);

  var card  = document.querySelector('.cmp-card');
  var modes = Array.prototype.slice.call(document.querySelectorAll('.mode'));
  modes.forEach(function(b){{
    b.addEventListener('click', function(){{
      card.classList.toggle('area', b.dataset.mode === 'area');
      modes.forEach(function(m){{ m.setAttribute('aria-pressed', String(m === b)); }});
    }});
  }});

  var all = Array.prototype.slice.call(document.querySelectorAll('details.act'));
  var btn = document.getElementById('expand');
  function setAll(open){{
    all.forEach(function(d){{ d.open = open; }});
    btn.setAttribute('aria-pressed', String(open));
    btn.textContent = open ? 'Close all sub-tasks' : 'Open all sub-tasks';
  }}
  btn.addEventListener('click', function(){{ setAll(btn.getAttribute('aria-pressed') !== 'true'); }});
  window.addEventListener('beforeprint', function(){{ setAll(true); }});
}})();
</script>
</body>
</html>
"""


def main():
    everyone = read_units()
    for site in SITES:
        units = [u for u in everyone if u["site"] == site[0]]
        if not units:
            continue
        out = ROOT / page_file(site[0])
        out.write_text(render(site, units, everyone), encoding="utf-8")
        flagged = sum(1 for u in units for a in u["activities"]
                      for x in [a] + a["subs"] + ([a["other"]] if a["other"] else [])
                      if x["flags"]) + sum(len(u["total_flags"]) for u in units)
        print(f"wrote {out} ({out.stat().st_size:,} bytes) "
              f"for {len(units)} unit{'s' if len(units) > 1 else ''}: "
              f"{', '.join(u['name'] for u in units)}; flagged {flagged}")


if __name__ == "__main__":
    main()
